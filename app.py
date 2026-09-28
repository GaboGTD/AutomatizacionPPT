import io
import os
import re
from datetime import datetime
from typing import Dict, Any, Tuple

import openpyxl
import pandas as pd
import streamlit as st
from pptx import Presentation
from pptx.util import Pt

# ==========================================
# CONFIGURACIÓN DE PÁGINA
# ==========================================
st.set_page_config(
    page_title="Generador de Presentaciones GTD",
    page_icon="📊",
    layout="centered"
)

TAMANO_FUENTE = Pt(9)
TABLA_GEN_CAL = 'DatosGenerales'
TABLA_QA_CAL = 'TablaCalificacion'
TABLA_GEN_RIESGO = 'Riesgo1'
TABLA_QA_RIESGO = 'Riesgo2'
RUTA_PLANTILLA = os.path.join(os.path.dirname(__file__), "Plantilla_Limpia.pptx")

# ==========================================
# FUNCIONES AUXILIARES
# ==========================================
def limpiar_texto(texto: Any) -> str:
    """Elimina espacios redundantes, saltos de línea y pasa a minúsculas."""
    if pd.isna(texto) or texto is None:
        return ""
    texto_limpio = re.sub(r'\s+', ' ', str(texto))
    return texto_limpio.strip().lower()

def aplicar_formato_texto(celda: Any, texto: Any) -> None:
    """Inserta texto y fuerza el tamaño de fuente fijo para anular el AutoFit de PPT."""
    celda.text = str(texto) if pd.notna(texto) else ''
    if celda.text and celda.text_frame.paragraphs:
        celda.text_frame.paragraphs[0].font.size = TAMANO_FUENTE

def encontrar_tabla(wb: openpyxl.Workbook, nombre_tabla: str):
    """Localiza una tabla con nombre en cualquier hoja del libro."""
    for sheet in wb.worksheets:
        if nombre_tabla in sheet.tables:
            return sheet, sheet.tables[nombre_tabla]
    return None, None

def extraer_tabla_kv(wb: openpyxl.Workbook, nombre_tabla: str) -> Dict[str, Any]:
    """Extrae datos de tablas horizontales (Fila 1: Nombres, Fila 2: Valores)."""
    sheet, tabla = encontrar_tabla(wb, nombre_tabla)
    if not sheet:
        return {}
    data_rows = [[cell.value for cell in row] for row in sheet[tabla.ref]]
    datos = {}
    if len(data_rows) > 1:
        keys, values = data_rows[0], data_rows[1]
        for i, key in enumerate(keys):
            if key and i < len(values) and pd.notna(values[i]):
                datos[limpiar_texto(key)] = values[i]
                datos[str(key).strip()] = values[i]
    return datos

def extraer_tabla_qa(wb: openpyxl.Workbook, nombre_tabla: str) -> pd.DataFrame:
    """Extrae tablas de declaraciones/preguntas."""
    sheet, tabla = encontrar_tabla(wb, nombre_tabla)
    if not sheet:
        return pd.DataFrame()
    data_rows = [[cell.value for cell in row] for row in sheet[tabla.ref]]
    if len(data_rows) > 1:
        return pd.DataFrame(data_rows[1:], columns=data_rows[0])
    return pd.DataFrame()

def preparar_dict_qa(df: pd.DataFrame) -> Dict[str, pd.Series]:
    """Indexa un DataFrame por su primera columna limpia."""
    if df.empty:
        return {}
    col_clave = df.columns[0]
    diccionario = {}
    for _, fila in df.iterrows():
        llave = limpiar_texto(fila[col_clave])
        if llave:
            diccionario[llave] = fila
    return diccionario

def obtener_valor_campo(fila_excel: pd.Series, patron: str) -> str:
    """Busca en la fila de Excel una columna cuyo nombre coincida con el patrón."""
    for col in fila_excel.index:
        col_clean = limpiar_texto(col)
        if patron in col_clean:
            val = fila_excel[col]
            return str(val) if pd.notna(val) else ''
    return ''

# ==========================================
# MOTOR DE GENERACIÓN
# ==========================================
def procesar_presentacion(excel_bytes: bytes) -> Tuple[io.BytesIO, str, Dict[str, Any]]:
    excel_stream = io.BytesIO(excel_bytes)
    wb = openpyxl.load_workbook(excel_stream, data_only=True)

    # 1. Extracción de tablas generales
    datos_gen_cal = extraer_tabla_kv(wb, TABLA_GEN_CAL)
    datos_gen_riesgo = extraer_tabla_kv(wb, TABLA_GEN_RIESGO)

    if not datos_gen_riesgo:
        datos_gen_riesgo = datos_gen_cal

    if not datos_gen_cal:
        raise ValueError(f"No se encontró la tabla '{TABLA_GEN_CAL}' en el archivo Excel.")

    # 2. Extracción de tablas de preguntas
    df_cal = extraer_tabla_qa(wb, TABLA_QA_CAL)
    df_riesgo = extraer_tabla_qa(wb, TABLA_QA_RIESGO)
    dict_cal = preparar_dict_qa(df_cal)
    dict_riesgo = preparar_dict_qa(df_riesgo)

    if not os.path.exists(RUTA_PLANTILLA):
        raise FileNotFoundError(f"No se encontró la plantilla en: {RUTA_PLANTILLA}")

    prs = Presentation(RUTA_PLANTILLA)

    # 3. Procesamiento de diapositivas
    for slide in prs.slides:
        es_slide_riesgo = False
        for shape in slide.shapes:
            if shape.has_text_frame and "riesgo" in shape.text.lower():
                es_slide_riesgo = True
                break

        datos_kv_actual = datos_gen_riesgo if es_slide_riesgo else datos_gen_cal

        for shape in slide.shapes:
            if not shape.has_table:
                continue
            tabla = shape.table

            # Identificar preguntas presentes en la tabla de PPT
            preguntas_ppt = []
            for i, row in enumerate(tabla.rows):
                if i > 0:
                    t = limpiar_texto(row.cells[0].text)
                    if t:
                        preguntas_ppt.append(t)

            matches_cal = sum(1 for p in preguntas_ppt if p in dict_cal)
            matches_riesgo = sum(1 for p in preguntas_ppt if p in dict_riesgo)

            # Si es tabla grande de preguntas
            if matches_cal > 0 or matches_riesgo > 0:
                tipo = "CALIFICACIÓN" if matches_cal >= matches_riesgo else "RIESGO"
                datos_excel = dict_cal if tipo == "CALIFICACIÓN" else dict_riesgo

                for i, row in enumerate(tabla.rows):
                    if i == 0:
                        continue
                    clave = limpiar_texto(row.cells[0].text)
                    if clave in datos_excel:
                        fila = datos_excel[clave]

                        # 1. Rellenar columna 'Respuesta' (columna 1 de PPT)
                        respuesta = obtener_valor_campo(fila, 'respuesta')
                        if len(row.cells) > 1:
                            aplicar_formato_texto(row.cells[1], respuesta)

                        # 2. Rellenar columna 'Comentario/Mitigación' (columna 2 de PPT)
                        # Busca por 'comentario' o por 'mitiga' para evitar problemas con columnas numéricas intermedias
                        comentario = obtener_valor_campo(fila, 'comentario')
                        if not comentario:
                            comentario = obtener_valor_campo(fila, 'mitiga')

                        if len(row.cells) > 2:
                            aplicar_formato_texto(row.cells[2], comentario)
            else:
                # Tabla superior (Clave-Valor)
                for row in tabla.rows:
                    for c, cell in enumerate(row.cells):
                        label = limpiar_texto(cell.text)
                        if label in datos_kv_actual and c + 1 < len(row.cells):
                            aplicar_formato_texto(row.cells[c + 1], datos_kv_actual[label])

    # 4. Construcción del nombre del archivo
    nombre_op = "Oportunidad_Sin_Nombre"
    for k, v in datos_gen_cal.items():
        if 'nombre' in k.lower() and 'oportunidad' in k.lower():
            nombre_op = str(v)
            break

    nombre_limpio = "".join(c for c in nombre_op if c.isalnum() or c in (' ', '_')).rstrip()
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    nombre_archivo = f"Calificación_{nombre_limpio}_{timestamp}.pptx"

    ppt_buffer = io.BytesIO()
    prs.save(ppt_buffer)
    ppt_buffer.seek(0)
    return ppt_buffer, nombre_archivo, datos_gen_cal

# ==========================================
# INTERFAZ WEB (STREAMLIT)
# ==========================================
st.title("Generador Automático de Presentaciones GTD")
st.markdown("Sube tu matriz de calificación y riesgo actualizada en Excel para compilar el PPT corporativo.")

archivo_subido = st.file_uploader(
    "Selecciona el archivo Excel",
    type=["xlsx", "xlsm"],
    help="Recuerda guardar (Ctrl + G) tu Excel antes de subirlo."
)

if archivo_subido is not None:
    contenido_excel = archivo_subido.getvalue()

    with st.spinner("Procesando Excel y rellenando PowerPoint..."):
        try:
            buffer, nombre_salida, info_detectada = procesar_presentacion(contenido_excel)
            
            st.success("¡Presentación generada correctamente!")
            
            st.caption(f"📌 **Datos detectados:** Cliente: `{info_detectada.get('cliente', 'N/A')}` | Oportunidad: `{info_detectada.get('nombre_oportunidad', info_detectada.get('Nombre Oportunidad (Código NCC)', 'N/A'))}`")

            st.download_button(
                label="📥 Descargar Presentación PPTX",
                data=buffer,
                file_name=nombre_salida,
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                type="primary"
            )
        except Exception as e:
            st.error(f"Error al procesar el archivo: {e}")
